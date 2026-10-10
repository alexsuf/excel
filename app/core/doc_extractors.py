import base64
import io
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple


ODD_SPACES = "\u00A0\u202F\u2009\u2007\u2005\u2006\u2008\u200A"
_SPACE_TRANSLATION = {ord(ch): " " for ch in ODD_SPACES}
_SPACE_RUN_RE = re.compile(r"[ \t]{2,}")
_BLANK_RUN_RE = re.compile(r"\n{3,}")
_HYPHEN_BREAK_RE = re.compile(r"(\w)[-\u2010\u2011]\s*\n\s*(\w)")
_DECIMAL_BREAK_RE = re.compile(r"(\d)[ \u00A0]*,[ \u00A0]*\n\s*(\d{2})\b")
_TRAILING_WS_RE = re.compile(r"[ \t]+$", re.MULTILINE)

MIN_PAGE_TEXT_CHARS = 40
OCR_GAIN_RATIO = 2.5
MAX_TEXT_PER_DOC = 200_000

EXT_TO_MIME = {
    '.pdf': 'application/pdf',
    '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    '.doc': 'application/msword',
    '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    '.xls': 'application/vnd.ms-excel',
    '.csv': 'text/csv',
    '.txt': 'text/plain',
    '.md': 'text/markdown',
    '.json': 'application/json',
    '.png': 'image/png',
    '.jpg': 'image/jpeg',
    '.jpeg': 'image/jpeg',
    '.tif': 'image/tiff',
    '.tiff': 'image/tiff',
    '.bmp': 'image/bmp',
}


def _log(msg: str):
    sys.stderr.write(f"[docx-extract] {msg}\n")
    sys.stderr.flush()


def _decode_safely(data: Any) -> str:
    if isinstance(data, bytes):
        return data.decode('utf-8', errors='replace')
    if isinstance(data, str):
        return data
    return str(data)


def _load_pymupdf():
    try:
        import pymupdf
        return pymupdf
    except ImportError:
        pass
    try:
        import fitz
        return fitz
    except ImportError:
        return None


def _load_ocr():
    try:
        import pytesseract
    except ImportError:
        pytesseract = None
    try:
        from PIL import Image
    except ImportError:
        Image = None
    return pytesseract, Image


def clean_text(text: str) -> str:
    if not text:
        return ""
    text = text.translate(_SPACE_TRANSLATION)
    text = _HYPHEN_BREAK_RE.sub(r"\1\2", text)
    text = _DECIMAL_BREAK_RE.sub(r"\1,\2", text)
    text = _TRAILING_WS_RE.sub("", text)
    text = _SPACE_RUN_RE.sub("  ", text)
    text = _BLANK_RUN_RE.sub("\n\n", text)
    return text.strip()


@dataclass
class DocumentExtraction:
    file_path: str
    display_name: str
    content: str = ""
    success: bool = True
    metadata: Dict[str, Any] = field(default_factory=dict)
    error: Optional[str] = None

    @property
    def is_empty(self) -> bool:
        return not self.content.strip()

    @property
    def needs_vision(self) -> bool:
        return bool(self.metadata.get("needs_vision"))


class DocumentTextExtractor:
    def __init__(self, logger=_log, ocr_enabled: bool = True, ocr_langs: str = "rus+eng",
                 ocr_zoom: float = 2.5):
        self.log = logger
        self.ocr_enabled = ocr_enabled
        self.ocr_langs = ocr_langs
        self.ocr_zoom = ocr_zoom
        self._pymupdf = _load_pymupdf()
        if ocr_enabled:
            self._tesseract, self._pil = _load_ocr()
        else:
            self._tesseract = None
            self._pil = _load_ocr()[1]

    @property
    def ocr_available(self) -> bool:
        return bool(self._tesseract and self._pil and self._pymupdf)

    def extract(self, file_path: str, original_name: Optional[str] = None) -> DocumentExtraction:
        path = Path(file_path)
        display_name = original_name or path.name
        if not path.exists():
            return DocumentExtraction(str(path), display_name, success=False,
                                      error="File not found")
        ext = path.suffix.lower()
        try:
            if ext == '.pdf':
                return self._extract_pdf(path, display_name)
            if ext in ('.docx', '.doc'):
                return self._extract_docx(path, display_name)
            if ext in ('.xlsx', '.xls'):
                return self._extract_excel(path, display_name)
            if ext in ('.png', '.jpg', '.jpeg', '.tif', '.tiff', '.bmp'):
                return self._extract_image(path, display_name)
            return self._extract_plain(path, display_name, ext)
        except Exception as exc:
            err = _decode_safely(exc)
            self.log(f"FAILED {display_name}: {err}")
            return DocumentExtraction(str(path), display_name, success=False, error=err)

    def _extract_pdf(self, path: Path, display_name: str) -> DocumentExtraction:
        if not self._pymupdf:
            return DocumentExtraction(str(path), display_name, success=False,
                                      error="PyMuPDF is not installed")
        fitz = self._pymupdf
        doc = fitz.open(str(path))
        page_count = len(doc)
        all_text_parts: List[str] = []
        ocr_pages = 0
        text_pages = 0
        empty_pages: List[int] = []
        native_texts = []
        try:
            for pno, page in enumerate(doc):
                page_no = pno + 1
                try:
                    native = _decode_safely(page.get_text("text", sort=True))
                except Exception:
                    native = ""

                native_len = len(native.strip())

                if native_len >= MIN_PAGE_TEXT_CHARS:
                    text_pages += 1
                    all_text_parts.append(f"--- {display_name}, стр. {page_no} (text) ---\n{native.strip()}")
                    native_texts.append(native)
                else:
                    if self.ocr_available and self._page_needs_ocr(page, native_len):
                        ocr_text = self._ocr_page(page)
                        ocr_len = len(ocr_text.strip())
                        if ocr_text.strip() and (ocr_len > OCR_GAIN_RATIO * max(native_len, 1) or native_len == 0):
                            ocr_pages += 1
                            all_text_parts.append(f"--- {display_name}, стр. {page_no} (OCR) ---\n{ocr_text.strip()}")
                        else:
                            empty_pages.append(page_no)
                    else:
                        empty_pages.append(page_no)

            ocr_text = "\n".join([p for p in all_text_parts if "(OCR)" in p])
            if native_texts:
                native_text = "\n".join(native_texts)
            else:
                native_text = ""
            
            table_text = self._extract_tables_from_text(native_text, ocr_text, display_name)
            
            if table_text:
                all_text_parts.append(table_text)
        finally:
            doc.close()

        content = clean_text("\n\n".join(all_text_parts))
        needs_vision = bool(empty_pages) and not ocr_pages and (not content or content.count("---") <= len(empty_pages))
        if empty_pages:
            self.log(f"{display_name}: no text on pages {empty_pages} (vision fallback needed)")

        return DocumentExtraction(
            file_path=str(path),
            display_name=display_name,
            content=content[:MAX_TEXT_PER_DOC],
            success=bool(content),
            metadata={
                "pages": page_count,
                "text_pages": text_pages,
                "ocr_pages": ocr_pages,
                "empty_pages": empty_pages,
                "needs_vision": needs_vision,
                "extraction": "text+ocr+tables" if (ocr_pages or table_text) else "text" if text_pages else "empty",
            },
            error=None if content else "PDF has no extractable text layer (scanned document)",
        )

    def _page_needs_ocr(self, page, native_len: int) -> bool:
        if native_len >= MIN_PAGE_TEXT_CHARS:
            return False
        try:
            return bool(page.get_images(full=True))
        except Exception:
            return True

    def _ocr_page(self, page) -> str:
        fitz = self._pymupdf
        try:
            pix = page.get_pixmap(matrix=fitz.Matrix(self.ocr_zoom, self.ocr_zoom))
            image = self._pil.open(io.BytesIO(pix.tobytes("png")))
            text = self._tesseract.image_to_string(image, lang=self.ocr_langs)
            return _decode_safely(text)
        except Exception as exc:
            self.log(f"OCR failed: {exc}")
            return ""

    def _extract_tables_from_text(self, native_text: str, ocr_text: str, display_name: str) -> str:
        blocks = []
        
        if native_text:
            for line in native_text.split("\n"):
                line = line.strip()
                if "| " in line or ("  " in line and len(line.split()) >= 4):
                    if "|" in line:
                        blocks.append(line)
                    else:
                        normalized = " | ".join(line.split())
                        if normalized.count(" | ") >= 5:
                            blocks.append(normalized)
        
        ocr_sections = {}
        current_section = []
        current_page = 0
        for line in ocr_text.split("\n"):
            if "стр." in line and "(OCR)" in line:
                if current_section:
                    ocr_sections[current_page] = "\n".join(current_section)
                current_section = []
                current_page += 1
                continue
            current_section.append(line)
        if current_section:
            ocr_sections[current_page] = "\n".join(current_section)
        
        for page, text in ocr_sections.items():
            table_block = self._parse_table_from_text(text)
            if table_block:
                blocks.append(table_block)
        
        return "\n\n".join(blocks)

    def _parse_table_from_text(self, text: str) -> str:
        if not text or len(text.strip()) < 30:
            return ""
        
        lines = text.strip().split("\n")
        table_lines = []
        in_table = False
        table_start_indicators = ["№", "Номер", "№ ", "товар", "наименование", "|"]
        
        for line in lines:
            stripped = line.strip()
            if not stripped:
                if in_table:
                    break
                continue
            
            is_header = any(ind.lower() in stripped.lower() for ind in ["№", "наименование", "кол-во", "цена", "сумма"])
            has_pipes = "|" in stripped
            looks_like_row = has_pipes and stripped.count("|") >= 5
            
            if is_header or (not in_table and (has_pipes or looks_like_row)):
                in_table = True
            
            if in_table:
                if has_pipes:
                    table_lines.append(stripped)
                elif stripped.count("  ") >= 4:
                    normalized = " | ".join(stripped.split())
                    if normalized.count("|") >= 4:
                        table_lines.append(normalized)
        
        if len(table_lines) >= 2:
            return "--- Таблица (из OCR) ---\n" + "\n".join(table_lines)
        return ""

    def _extract_docx(self, path: Path, display_name: str) -> DocumentExtraction:
        try:
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError:
            return DocumentExtraction(str(path), display_name, success=False,
                                      error="python-docx is not installed")
        try:
            doc = Document(str(path))
        except Exception as exc:
            err = _decode_safely(exc)
            hint = "corrupted" if ("package" in err.lower() or "corrupt" in err.lower()) else "could not be opened"
            return DocumentExtraction(str(path), display_name, success=False,
                                      error=f"DOCX {hint}: {err}")

        parts: List[str] = []
        body = doc.element.body
        for child in body.iterchildren():
            tag = child.tag.split('}')[-1]
            if tag == 'p':
                text = clean_text(Paragraph(child, doc).text)
                if text:
                    parts.append(text)
            elif tag == 'tbl':
                table = Table(child, doc)
                for row in table.rows:
                    cells = [clean_text(cell.text) for cell in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))

        content = clean_text("\n".join(parts))
        return DocumentExtraction(
            file_path=str(path),
            display_name=display_name,
            content=content[:MAX_TEXT_PER_DOC],
            success=bool(content),
            metadata={"paragraphs": len(doc.paragraphs), "tables": len(doc.tables),
                      "extraction": "docx"},
            error=None if content else "DOCX contains no text",
        )

    def _extract_excel(self, path: Path, display_name: str) -> DocumentExtraction:
        ext = path.suffix.lower()
        parts: List[str] = []
        sheets = 0
        rows_total = 0
        try:
            if ext == '.xls':
                import xlrd
                book = xlrd.open_workbook(str(path), encoding_override='utf-8')
                for name in book.sheet_names():
                    sheet = book.sheet_by_name(name)
                    parts.append(f"[Лист: {name}]")
                    for r in range(sheet.nrows):
                        cells = [_decode_safely(v) for v in sheet.row_values(r)]
                        line = " | ".join(c.strip() for c in cells)
                        if line.strip(" |"):
                            parts.append(line)
                            rows_total += 1
                    sheets += 1
            else:
                import openpyxl
                book = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
                try:
                    for name in book.sheetnames:
                        sheet = book[name]
                        parts.append(f"[Лист: {name}]")
                        for row in sheet.iter_rows(values_only=True):
                            cells = ["" if c is None else _decode_safely(c) for c in row]
                            line = " | ".join(c.strip() for c in cells)
                            if line.strip(" |"):
                                parts.append(line)
                                rows_total += 1
                        sheets += 1
                finally:
                    book.close()
        except ImportError:
            return DocumentExtraction(str(path), display_name, success=False,
                                      error="openpyxl/xlrd is not installed")
        except Exception as exc:
            err = _decode_safely(exc)
            low = err.lower()
            if "password" in low or "encrypted" in low:
                return DocumentExtraction(str(path), display_name, success=False,
                                          error=f"Excel file is password-protected: {err}")
            return DocumentExtraction(str(path), display_name, success=False,
                                      error=f"Failed to read Excel file: {err}")

        content = clean_text("\n".join(parts))
        return DocumentExtraction(
            file_path=str(path),
            display_name=display_name,
            content=content[:MAX_TEXT_PER_DOC],
            success=bool(content),
            metadata={"sheets": sheets, "rows": rows_total, "extension": ext,
                      "extraction": "excel"},
            error=None if content else "Excel file contains no data",
        )

    def _extract_image(self, path: Path, display_name: str) -> DocumentExtraction:
        if not self.ocr_available:
            return DocumentExtraction(
                str(path), display_name,
                content=f"[Изображение: {display_name}]",
                success=False,
                metadata={"needs_vision": True},
                error="OCR is unavailable (install pytesseract, pillow and tesseract-ocr)",
            )
        try:
            image = self._pil.open(str(path))
            text = _decode_safely(self._tesseract.image_to_string(image, lang=self.ocr_langs))
        except Exception as exc:
            err = _decode_safely(exc)
            return DocumentExtraction(str(path), display_name, success=False,
                                      metadata={"needs_vision": True},
                                      error=f"OCR failed: {err}")
        content = clean_text(text)
        return DocumentExtraction(
            file_path=str(path),
            display_name=display_name,
            content=content[:MAX_TEXT_PER_DOC],
            success=bool(content),
            metadata={"extraction": "ocr"},
            error=None if content else "OCR returned no text",
        )

    def _extract_plain(self, path: Path, display_name: str, ext: str) -> DocumentExtraction:
        text = ""
        for encoding in ('utf-8', 'utf-8-sig', 'cp1251'):
            try:
                with open(path, 'r', encoding=encoding) as handle:
                    text = handle.read()
                break
            except UnicodeDecodeError:
                continue
            except Exception as exc:
                err = _decode_safely(exc)
                return DocumentExtraction(str(path), display_name, success=False,
                                          error=f"Failed to read text file: {err}")
        content = clean_text(text)
        return DocumentExtraction(
            file_path=str(path),
            display_name=display_name,
            content=content[:MAX_TEXT_PER_DOC],
            success=bool(content),
            metadata={"extension": ext, "extraction": "text"},
            error=None if content else "Text file is empty",
        )

    def render_pages_as_images(self, file_path: str, max_pages: int = 20,
                               zoom: float = 1.6, jpeg_quality: int = 80) -> List[Tuple[int, str, str]]:
        fitz = self._pymupdf
        if not fitz:
            return []
        doc = fitz.open(str(file_path))
        out: List[Tuple[int, str, str]] = []
        try:
            for pno, page in enumerate(doc):
                if len(out) >= max_pages:
                    break
                try:
                    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
                    encoded = self._encode_pixmap(pix, jpeg_quality)
                except Exception as exc:
                    self.log(f"page {pno + 1} render failed: {exc}")
                    continue
                if encoded:
                    mime, b64 = encoded
                    out.append((pno + 1, mime, b64))
        finally:
            doc.close()
        return out

    def _encode_pixmap(self, pix, jpeg_quality: int):
        if self._pil:
            try:
                mode = "RGB" if pix.n >= 3 else "L"
                image = self._pil.frombytes(mode, (pix.width, pix.height), pix.samples)
                buffer = io.BytesIO()
                image.save(buffer, format="JPEG", quality=jpeg_quality, optimize=True)
                return "image/jpeg", base64.b64encode(buffer.getvalue()).decode("ascii")
            except Exception:
                pass
        try:
            return "image/png", base64.b64encode(pix.tobytes("png")).decode("ascii")
        except Exception:
            return None


_extractor_singleton: Optional[DocumentTextExtractor] = None


def get_document_extractor() -> DocumentTextExtractor:
    global _extractor_singleton
    if _extractor_singleton is None:
        from app.core.config import settings
        _extractor_singleton = DocumentTextExtractor(
            ocr_enabled=getattr(settings, "ocr_enabled", True),
            ocr_langs=getattr(settings, "ocr_langs", "rus+eng"),
            ocr_zoom=getattr(settings, "ocr_zoom", 2.5),
        )
    return _extractor_singleton