import re
import sys
import time
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string
from typing import Dict, List, Optional, Any
from dataclasses import dataclass, field
import json
import os


def _debug_log(msg: str):
    ts = time.strftime('%H:%M:%S')
    line = f"[ta][{ts}] {msg}"
    sys.stderr.write(line + "\n")
    sys.stderr.flush()
    try:
        from app.core.config import settings
        log_path = Path(settings.output_dir) / "llm.log"
        log_path.parent.mkdir(parents=True, exist_ok=True)
        with open(log_path, "a", encoding="utf-8") as lf:
            lf.write(line + "\n")
            lf.flush()
            import os
            os.fsync(lf.fileno())
    except Exception:
        pass


@dataclass
class NamedRange:
    name: str
    sheet_name: str
    cell_range: str
    min_row: int
    max_row: int
    min_col: int
    max_col: int
    is_table: bool = False
    table_headers: List[str] = field(default_factory=list)
    value: Any = None
    data_type: str = "string"
    is_bookmark: bool = False


class TemplateAnalyzer:
    def __init__(self, template_path: str):
        self.template_path = template_path
        _debug_log(f"INIT START: loading workbook from {template_path}")
        t0 = time.time()
        self.workbook = openpyxl.load_workbook(template_path, data_only=True)
        _debug_log(f"INIT DONE: workbook loaded in {time.time()-t0:.2f}s, sheets={self.workbook.sheetnames}")
        self.named_ranges: Dict[str, NamedRange] = {}
        t1 = time.time()
        self._analyze_named_ranges()
        _debug_log(f"_analyze_named_ranges DONE in {time.time()-t1:.2f}s, found {len(self.named_ranges)} ranges")
        t2 = time.time()
        self._scan_bookmarks()
        _debug_log(f"_scan_bookmarks DONE in {time.time()-t2:.2f}s, total ranges={len(self.named_ranges)}")

    def _analyze_named_ranges(self):
        _debug_log("_analyze_named_ranges START")
        try:
            defined_name_iter = self.workbook.defined_names.definedNameList
            _debug_log(f"Got definedNameList, count={len(defined_name_iter)}")
        except AttributeError:
            defined_name_iter = [self.workbook.defined_names[n] for n in self.workbook.defined_names]
            _debug_log(f"Got defined_names via indexing, count={len(defined_name_iter)}")
        for idx, defined_name in enumerate(defined_name_iter):
            name = defined_name.name
            _debug_log(f"  Processing named range[{idx}]: {name}")
            t_start = time.time()
            try:
                destinations = list(defined_name.destinations)
                if not destinations:
                    _debug_log(f"    No destinations for '{name}', skipping")
                    continue
                sheet_name, cell_range = destinations[0]
                sheet_name = sheet_name.strip("'")
                _debug_log(f"    Range '{name}': sheet='{sheet_name}' range='{cell_range}'")
                if sheet_name not in self.workbook.sheetnames:
                    _debug_log(f"    Sheet '{sheet_name}' not in workbook sheets, skipping")
                    continue
                ws = self.workbook[sheet_name]
                min_row, min_col, max_row, max_col = self._parse_range(cell_range)
                _debug_log(f"    Parsed range: rows {min_row}-{max_row}, cols {min_col}-{max_col}")
                t_detect = time.time()
                is_table = self._detect_table(ws, min_row, min_col, max_row, max_col)
                _debug_log(f"    Table detection in {time.time()-t_detect:.3f}s: {is_table}")
                table_headers = []
                if is_table:
                    t_headers = time.time()
                    table_headers = self._extract_table_headers(ws, min_row, min_col, max_row, max_col)
                    _debug_log(f"    Table headers extracted in {time.time()-t_headers:.3f}s: {table_headers}")
                nr = NamedRange(
                    name=name,
                    sheet_name=sheet_name,
                    cell_range=cell_range,
                    min_row=min_row,
                    max_row=max_row,
                    min_col=min_col,
                    max_col=max_col,
                    is_table=is_table,
                    table_headers=table_headers
                )
                self.named_ranges[name] = nr
                _debug_log(f"    Added range '{name}' in {time.time()-t_start:.3f}s")
            except Exception as e:
                _debug_log(f"    ERROR processing range '{name}': {e}")
                continue
        _debug_log(f"_analyze_named_ranges END, total={len(self.named_ranges)}")

    def _scan_bookmarks(self):
        _debug_log("_scan_bookmarks START")
        bookmark_pattern = re.compile(r'^\s*\{\{(\w+)\}\}\s*$')
        total_cells_scanned = 0
        for sheet_name in self.workbook.sheetnames:
            ws = self.workbook[sheet_name]
            _debug_log(f"  Scanning sheet '{sheet_name}' for bookmarks...")
            for row_cells in ws.iter_rows():
                for cell in row_cells:
                    total_cells_scanned += 1
                    if not cell.value or not isinstance(cell.value, str):
                        continue
                    match = bookmark_pattern.match(cell.value)
                    if not match:
                        continue
                    bookmark_name = match.group(1)
                    _debug_log(f"    Found bookmark '{{{{{bookmark_name}}}}}'' at {cell.coordinate}")
                    if bookmark_name in self.named_ranges:
                        _debug_log(f"      Already exists as named range, skipping")
                        continue
                    t_bm = time.time()
                    self._create_bookmark_range(bookmark_name, sheet_name, cell.row, cell.column, ws)
                    _debug_log(f"      Bookmark '{bookmark_name}' created in {time.time()-t_bm:.3f}s")
        _debug_log(f"_scan_bookmarks END, scanned={total_cells_scanned}, total_ranges={len(self.named_ranges)}")

    def _create_bookmark_range(self, name: str, sheet_name: str, row: int, col: int, ws):
        col_idx = col + 1
        headers = []
        while True:
            cell = ws.cell(row=row, column=col_idx)
            if cell.value is None or str(cell.value).strip() == '':
                break
            headers.append(str(cell.value).strip())
            col_idx += 1
        if len(headers) >= 2:
            min_c = col + 1
            max_c = col + len(headers)
            cell_range = f"${get_column_letter(min_c)}${row}:${get_column_letter(max_c)}${row}"
            nr = NamedRange(
                name=name, sheet_name=sheet_name, cell_range=cell_range,
                min_row=row, max_row=row, min_col=min_c, max_c=max_c,
                is_table=True, table_headers=headers, is_bookmark=True
            )
            _debug_log(f"    Created table bookmark '{name}': {cell_range}, headers={headers}")
        else:
            cell_range = f"${get_column_letter(col)}${row}"
            nr = NamedRange(
                name=name, sheet_name=sheet_name, cell_range=cell_range,
                min_row=row, max_row=row, min_col=col, max_col=col,
                is_table=False, is_bookmark=True
            )
            _debug_log(f"    Created single bookmark '{name}': {cell_range}")
        self.named_ranges[name] = nr

    def _parse_range(self, cell_range: str) -> tuple:
        if ':' in cell_range:
            start, end = cell_range.split(':')
        else:
            start = end = cell_range
        start = start.replace('$', '')
        end = end.replace('$', '')
        start_col_str = ''.join(filter(str.isalpha, start))
        start_row_str = ''.join(filter(str.isdigit, start))
        end_col_str = ''.join(filter(str.isalpha, end))
        end_row_str = ''.join(filter(str.isdigit, end))
        min_col = column_index_from_string(start_col_str) if start_col_str else 1
        min_row = int(start_row_str) if start_row_str else 1
        max_col = column_index_from_string(end_col_str) if end_col_str else min_col
        max_row = int(end_row_str) if end_row_str else min_row
        return min_row, min_col, max_row, max_col

    def _detect_table(self, ws, min_row: int, min_col: int, max_row: int, max_col: int) -> bool:
        if max_row - min_row < 1 or max_col - min_col < 1:
            return False
        header_row = min_row
        header_values = []
        for c in range(min_col, max_col + 1):
            cell = ws.cell(row=header_row, column=c)
            if cell.value is not None:
                header_values.append(str(cell.value).strip())
        return len(header_values) > 0 and max_row > min_row

    def _extract_table_headers(self, ws, min_row: int, min_col: int, max_row: int, max_col: int) -> List[str]:
        headers = []
        for c in range(min_col, max_col + 1):
            cell = ws.cell(row=min_row, column=c)
            headers.append(str(cell.value).strip() if cell.value else f"Column_{c}")
        return headers

    def get_named_ranges(self) -> Dict[str, NamedRange]:
        return self.named_ranges

    def get_field_schema(self, dictionary_manager=None) -> Dict[str, Any]:
        _debug_log("get_field_schema START")
        t0 = time.time()
        schema = {"type": "object", "properties": {}, "required": []}
        for name, nr in self.named_ranges.items():
            mapping = dictionary_manager.get_mapping(name) if dictionary_manager else None
            if nr.is_table:
                col_props = {}
                for h in nr.table_headers:
                    col_mapping = self._get_table_column_mapping(name, h, mapping)
                    col_props[h] = {"type": col_mapping.get("type", "string")}
                schema["properties"][name] = {
                    "type": "array",
                    "items": {"type": "object", "properties": col_props},
                    "description": f"Table with columns: {', '.join(nr.table_headers)}"
                }
            else:
                data_type = mapping.data_type if mapping else "string"
                json_type = self._py_type_to_json_type(data_type)
                schema["properties"][name] = {
                    "type": json_type,
                    "description": (mapping.description if mapping else
                                    f"Cell range: {nr.cell_range} on sheet {nr.sheet_name}")
                }
                if mapping and mapping.required:
                    schema["required"].append(name)
        _debug_log(f"get_field_schema DONE in {time.time()-t0:.2f}s, fields={len(schema['properties'])}")
        return schema

    def _get_table_column_mapping(self, table_name: str, column: str, mapping) -> Dict[str, str]:
        if mapping and mapping.validation_rules:
            for rule, rule_val in mapping.validation_rules.items():
                if isinstance(rule_val, dict) and column in rule_val:
                    return {"type": self._py_type_to_json_type(rule_val[column].get("type", "string"))}
            col_type = mapping.validation_rules.get(f"{column}_type")
            if col_type:
                return {"type": self._py_type_to_json_type(col_type)}
        return {"type": "string"}

    @staticmethod
    def _py_type_to_json_type(py_type: str) -> str:
        mapping = {
            "string": "string", "str": "string",
            "number": "number", "float": "number",
            "integer": "integer", "int": "integer",
            "boolean": "boolean", "bool": "boolean",
            "date": "string", "datetime": "string"
        }
        return mapping.get(py_type, "string")

    def get_field_labels(self) -> Dict[str, str]:
        labels: Dict[str, str] = {}
        for name, nr in self.named_ranges.items():
            if nr.sheet_name not in self.workbook.sheetnames:
                continue
            ws = self.workbook[nr.sheet_name]
            try:
                parts: List[str] = []
                for col in range(nr.min_col - 1, 0, -1):
                    cell = ws.cell(row=nr.min_row, column=col)
                    if cell.value is not None:
                        label = str(cell.value).strip()
                        if label:
                            parts.append(label)
                            break
                if parts:
                    labels[name] = parts[0]
            except Exception:
                continue
        return labels

    def to_dict(self) -> Dict[str, Any]:
        result = {}
        for name, nr in self.named_ranges.items():
            result[name] = {
                "name": nr.name, "sheet_name": nr.sheet_name,
                "cell_range": nr.cell_range,
                "min_row": nr.min_row, "max_row": nr.max_row,
                "min_col": nr.min_col, "max_col": nr.max_col,
                "is_table": nr.is_table,
                "table_headers": nr.table_headers,
                "data_type": nr.data_type,
                "is_bookmark": nr.is_bookmark
            }
        return result


def analyze_template(template_path: str) -> TemplateAnalyzer:
    return TemplateAnalyzer(template_path)
