import openpyxl
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill, numbers
from openpyxl.utils import get_column_letter
from typing import Dict, Any, List, Optional
import pandas as pd
from app.core.template_analyzer import TemplateAnalyzer, NamedRange
from app.core.validator import ExtractionResult, DynamicField, TableData
import copy


class ExcelWriter:
    def __init__(self, template_path: str):
        self.template_path = template_path
        self.workbook = openpyxl.load_workbook(template_path)
        self.analyzer = TemplateAnalyzer(template_path)
        self.named_ranges = self.analyzer.get_named_ranges()
    
    def write_data(self, extraction_result: ExtractionResult, output_path: str):
        for name, nr in self.named_ranges.items():
            if nr.is_table:
                table_data = extraction_result.tables.get(name)
                if table_data:
                    self._write_table(nr, table_data)
            else:
                field_data = extraction_result.fields.get(name)
                if field_data:
                    self._write_field(nr, field_data)
                else:
                    self._clear_range(nr)
        
        self.workbook.save(output_path)

    def _clear_range(self, nr: NamedRange):
        ws = self.workbook[nr.sheet_name]
        for row in range(nr.min_row, nr.max_row + 1):
            for col in range(nr.min_col, nr.max_col + 1):
                ws.cell(row=row, column=col).value = None
    
    def _write_field(self, nr: NamedRange, field: DynamicField):
        ws = self.workbook[nr.sheet_name]
        
        if nr.min_row == nr.max_row and nr.min_col == nr.max_col:
            cell = ws.cell(row=nr.min_row, column=nr.min_col)
            self._set_cell_value(cell, field)
        else:
            for row in range(nr.min_row, nr.max_row + 1):
                for col in range(nr.min_col, nr.max_col + 1):
                    cell = ws.cell(row=row, column=col)
                    self._set_cell_value(cell, field)
    
    def _set_cell_value(self, cell, field: DynamicField):
        value = field.value

        if value is None:
            cell.value = None
            return
        
        if isinstance(value, (int, float)):
            cell.value = value
            if isinstance(value, float) and value == int(value):
                cell.number_format = '#,##0'
            elif isinstance(value, float):
                cell.number_format = '#,##0.00'
        
        elif isinstance(value, bool):
            cell.value = value
        
        elif hasattr(value, 'isoformat'):
            cell.value = value
            cell.number_format = 'DD.MM.YYYY'
        
        else:
            cell.value = str(value)
        
        if not field.is_valid:
            cell.fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
            cell.font = Font(color="CC0000")
    
    def _write_table(self, nr: NamedRange, table: TableData):
        ws = self.workbook[nr.sheet_name]
        
        for i, row_data in enumerate(table.rows):
            row_idx = nr.min_row + 1 + i
            
            if row_idx > nr.max_row:
                ws.insert_rows(row_idx)
                nr.max_row += 1
            
            for j, header in enumerate(nr.table_headers):
                col_idx = nr.min_col + j
                cell = ws.cell(row=row_idx, column=col_idx)
                value = row_data.cells.get(header)

                self._set_cell_value_simple(cell, value, header, table)
                
                if not row_data.is_valid:
                    cell.fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
    
    def _set_cell_value_simple(self, cell, value: Any, header: str, table: TableData):
        if value is None:
            cell.value = None
            return
        if isinstance(value, (int, float)):
            cell.value = value
            if isinstance(value, float) and value == int(value):
                cell.number_format = '#,##0'
            elif isinstance(value, float):
                cell.number_format = '#,##0.00'
        elif isinstance(value, bool):
            cell.value = value
        elif hasattr(value, 'isoformat'):
            cell.value = value
            cell.number_format = 'DD.MM.YYYY'
        else:
            cell.value = str(value)
    
    def write_json_direct(self, data: Dict[str, Any], output_path: str):
        for name, nr in self.named_ranges.items():
            value = data.get(name)
            if value is None:
                continue
            
            if nr.is_table:
                self._write_table_direct(nr, value)
            else:
                self._write_field_direct(nr, value)
        
        self.workbook.save(output_path)
    
    def _write_field_direct(self, nr: NamedRange, value: Any):
        ws = self.workbook[nr.sheet_name]
        
        if nr.min_row == nr.max_row and nr.min_col == nr.max_col:
            cell = ws.cell(row=nr.min_row, column=nr.min_col)
            self._set_cell_value_direct(cell, value)
        else:
            for row in range(nr.min_row, nr.max_row + 1):
                for col in range(nr.min_col, nr.max_col + 1):
                    cell = ws.cell(row=row, column=col)
                    self._set_cell_value_direct(cell, value)
    
    def _set_cell_value_direct(self, cell, value: Any):
        if value is None:
            cell.value = None
            return
        
        if isinstance(value, (int, float)):
            cell.value = value
            if isinstance(value, float) and value == int(value):
                cell.number_format = '#,##0'
            elif isinstance(value, float):
                cell.number_format = '#,##0.00'
        elif isinstance(value, bool):
            cell.value = value
        elif hasattr(value, 'isoformat'):
            cell.value = value
            cell.number_format = 'DD.MM.YYYY'
        else:
            cell.value = str(value)
    
    def _write_table_direct(self, nr: NamedRange, table_data: List[Dict[str, Any]]):
        ws = self.workbook[nr.sheet_name]
        
        for i, row_data in enumerate(table_data):
            row_idx = nr.min_row + 1 + i
            
            if row_idx > nr.max_row:
                ws.insert_rows(row_idx)
                nr.max_row += 1
            
            for j, header in enumerate(nr.table_headers):
                col_idx = nr.min_col + j
                cell = ws.cell(row=row_idx, column=col_idx)
                value = row_data.get(header)
                
                if value is not None:
                    self._set_cell_value_direct(cell, value)
    
    def get_template_info(self) -> Dict[str, Any]:
        return self.analyzer.to_dict()
    
    def preview_mapping(self, data: Dict[str, Any]) -> Dict[str, Any]:
        preview = {}
        for name, nr in self.named_ranges.items():
            value = data.get(name)
            if nr.is_table:
                preview[name] = {
                    "type": "table",
                    "headers": nr.table_headers,
                    "rows_count": len(value) if isinstance(value, list) else 0,
                    "sample": value[:3] if isinstance(value, list) else None
                }
            else:
                preview[name] = {
                    "type": "field",
                    "cell_range": nr.cell_range,
                    "value": value
                }
        return preview